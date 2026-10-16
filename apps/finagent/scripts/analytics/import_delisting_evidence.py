#!/usr/bin/env python3
"""Import reviewed delisting facts from hash-bound exchange JSON/XLSX files.

The request manifest cannot supply evidence by itself. Every requested date and
identity must match the saved official response. Last trading dates and delisting
decision dates are not delisting dates. No historical membership is deleted.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from scripts.analytics.common import connect
from psycopg.types.json import Jsonb

SOURCE = "exchange:reviewed_delisting:v1"


def date_value(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    if re.fullmatch(r"\d{8}", text):
        text = f"{text[:4]}-{text[4:6]}-{text[6:]}"
    return dt.date.fromisoformat(text)


def source_rows(source):
    """Parse only the declared official endpoint, retaining raw row provenance."""
    url = urlparse(source["url"])
    params = parse_qs(url.query)
    kind = source["format"]
    valid = url.scheme == "https" and not url.username and not url.password
    if kind == "sse_delisted_json":
        valid = valid and url.hostname == "query.sse.com.cn" and url.path == "/commonQuery.do"
        valid = valid and params.get("sqlId") == ["COMMON_SSE_CP_GPJCTPZ_GPLB_GP_L"]
        valid = valid and params.get("COMPANY_STATUS") == ["3"]
    elif kind == "szse_delisted_xlsx":
        valid = valid and url.hostname == "www.szse.cn" and url.path == "/api/report/ShowReport"
        valid = valid and params.get("CATALOGID") == ["1793_ssgs"] and params.get("TABKEY") == ["tab2"]
        valid = valid and params.get("SHOWTYPE") == ["xlsx"]
    else:
        raise ValueError("unsupported delisting source format")
    if not valid:
        raise ValueError("not the matching official delisted-list endpoint")
    observed = dt.datetime.fromisoformat(source["observed_at"])
    if observed.tzinfo is None or observed > dt.datetime.now(dt.timezone.utc) + dt.timedelta(minutes=5):
        raise ValueError("invalid observed timestamp")
    path = Path(source["file"])
    raw = path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != source["sha256"]:
        raise ValueError("source file hash mismatch")
    if kind == "sse_delisted_json":
        rows = json.loads(raw)["result"]
        result = [dict(symbol=str(r["A_STOCK_CODE"]), name=r["COMPANY_ABBR"],
                       listed_date=r["LIST_DATE"], delisted_date=r["DELIST_DATE"],
                       state=str(r["STATE_CODE"]), source_row=r, exchange="SSE") for r in rows]
    else:
        from openpyxl import load_workbook
        with path.open("rb") as stream:
            workbook = load_workbook(stream, read_only=True, data_only=True)
            try:
                # SZSE exports can declare A1:A1 despite containing a full table.
                # Ignore worksheet dimension metadata, never silently truncate.
                workbook.active.reset_dimensions()
                values = workbook.active.iter_rows(values_only=True)
                headers = next(values)
                required = {"证券代码", "证券简称", "上市日期", "终止上市日期"}
                if not required.issubset(headers):
                    raise ValueError("unexpected delisted XLSX schema")
                result = []
                for values_row in values:
                    row = {k: v.date().isoformat() if isinstance(v, dt.datetime) else
                           v.isoformat() if isinstance(v, dt.date) else v
                           for k, v in zip(headers, values_row)}
                    if row["证券代码"] is None:
                        continue
                    result.append(dict(symbol=str(row["证券代码"]).zfill(6), name=row["证券简称"],
                                       listed_date=row["上市日期"], delisted_date=row["终止上市日期"],
                                       state="3", source_row=row, exchange="SZSE"))
            finally:
                workbook.close()
    return result, observed


def validate_manifest(manifest):
    if manifest.get("schema") != "reviewed-delisting-v1" or not manifest.get("items"):
        raise ValueError("empty or unsupported delisting manifest")
    cache, matched, identities = {}, [], set()
    for item in manifest["items"]:
        if not item.get("reviewed"):
            raise ValueError("reviewed evidence required")
        identity = item["exchange"], item["symbol"]
        if identity in identities:
            raise ValueError("duplicate requested identity")
        identities.add(identity)
        source = item["source"]
        key = json.dumps(source, sort_keys=True)
        if key not in cache:
            cache[key] = source_rows(source)
        rows, observed = cache[key]
        rows = [r for r in rows if (r["exchange"], r["symbol"]) == identity]
        if len(rows) != 1:
            raise ValueError("identity absent or ambiguous in official source: " + item["symbol"])
        row = rows[0]
        prefix = ("600", "601", "603", "605", "688", "689") if identity[0] == "SSE" else ("000", "001", "002", "003", "300", "301")
        if not re.fullmatch(r"\d{6}", row["symbol"]) or not row["symbol"].startswith(prefix):
            raise ValueError("not an A share identity")
        date, listed = date_value(row["delisted_date"]), date_value(row["listed_date"])
        if row["state"] != "3" or date != date_value(item["delisted_date"]):
            raise ValueError("requested delisting not supported by official source")
        if not dt.date(1980, 1, 1) <= listed < date <= observed.astimezone(ZoneInfo("Asia/Shanghai")).date():
            raise ValueError("invalid listing/delisting chronology")
        matched.append({**row, "listed_date": listed.isoformat(), "delisted_date": date.isoformat(),
                        "source": source, "reviewed": item["reviewed"]})
    return matched


def apply_rows(conn, rows, output_dir, *, apply=False):
    """One transaction; archive originals before writing any evidence or metadata."""
    conn.execute("SET LOCAL statement_timeout='25s'")
    conn.execute("SET LOCAL lock_timeout='5s'")
    if apply:
        conn.execute("SELECT pg_advisory_xact_lock(hashtextextended('reviewed-delisting-v1',0))")
    assets = conn.execute("""SELECT * FROM market.assets WHERE asset_type='stock'
        AND NOT(metadata ? 'identity_alias_of') AND symbol=ANY(%s) ORDER BY asset_id""",
        ([r["symbol"] for r in rows],)).fetchall()
    by_key = {}
    for asset in assets:
        key = asset["exchange"], asset["symbol"]
        if key in by_key:
            raise ValueError("ambiguous canonical identity: " + asset["symbol"])
        by_key[key] = asset
    for row in rows:
        asset = by_key.get((row["exchange"], row["symbol"]))
        if not asset:
            raise ValueError("canonical identity missing: " + row["symbol"])
        if asset["delisted_date"] and str(asset["delisted_date"]) != row["delisted_date"]:
            raise ValueError("existing delisting date conflicts: " + row["symbol"])
        if conn.execute("""SELECT 1 FROM market.bars_daily WHERE asset_id=%s
            AND ts >= (%s::date::timestamp AT TIME ZONE 'Asia/Shanghai') AND close>0 LIMIT 1""",
            (asset["asset_id"], row["delisted_date"])).fetchone():
            raise ValueError("price on/after declared delisting requires review: " + row["symbol"])
    prior = conn.execute("SELECT * FROM market.asset_lifecycle_versions WHERE asset_id=ANY(%s) ORDER BY lifecycle_version_id",
                         ([a["asset_id"] for a in assets],)).fetchall()
    output_dir.mkdir(parents=True, exist_ok=True)
    tag = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%S.%fZ")
    (output_dir / (tag + "-before.json")).write_text(json.dumps({"assets": assets, "lifecycle": prior}, default=str, ensure_ascii=False, indent=2))
    inserted = updated = 0
    if apply:
        for row in rows:
            asset = by_key[row["exchange"], row["symbol"]]
            evidence_hash = hashlib.sha256(json.dumps(row, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()
            evidence = {**row, "evidence_sha256": evidence_hash,
                        "verified_fields": ["delisted_date", "lifecycle_status", "is_tradable"],
                        "is_st_verified": False,
                        "knowledge_semantics": "available_at is actual import time; no historical knowledge backdating",
                        "scope": "official delisting effective on this date; no prior tradability or liquidation-price assertion"}
            inserted += conn.execute("""INSERT INTO market.asset_lifecycle_versions(
                asset_id,valid_from,valid_to,lifecycle_status,is_tradable,is_st,is_suspended,exchange,
                source_name,available_at,quality_status,metadata)
                SELECT %s,%s,NULL,'delisted',false,false,false,%s,%s,now(),'verified',%s
                WHERE NOT EXISTS(SELECT 1 FROM market.asset_lifecycle_versions WHERE asset_id=%s
                  AND source_name=%s AND metadata->>'evidence_sha256'=%s)""",
                (asset["asset_id"], row["delisted_date"], row["exchange"], SOURCE, Jsonb(evidence),
                 asset["asset_id"], SOURCE, evidence_hash)).rowcount
            metadata = {"delisting_date_verified": True, "delisting_date_source": row["source"],
                        "delisting_evidence_sha256": evidence_hash,
                        "lifecycle_scope": "current master state; use lifecycle versions with cutoff for historical queries"}
            updated += conn.execute("""UPDATE market.assets SET delisted_date=%s,lifecycle_status='delisted',
                metadata=metadata||%s,updated_at=now() WHERE asset_id=%s AND
                (delisted_date IS DISTINCT FROM %s::date OR lifecycle_status IS DISTINCT FROM 'delisted'
                 OR metadata->>'delisting_evidence_sha256' IS DISTINCT FROM %s)""",
                (row["delisted_date"], Jsonb(metadata), asset["asset_id"], row["delisted_date"], evidence_hash)).rowcount
    return {"reviewed_assets": len(rows), "applied": apply, "lifecycle_inserted": inserted,
            "master_updated": updated, "prices_modified": 0, "positions_removed": 0,
            "historical_system_knowledge_backdated": False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    rows = validate_manifest(json.loads(args.manifest.read_text()))
    with connect(autocommit=False) as conn:
        if not args.apply:
            conn.execute("SET TRANSACTION READ ONLY")
        result = apply_rows(conn, rows, args.output_dir, apply=args.apply)
    result["at_utc"] = dt.datetime.now(dt.timezone.utc).isoformat()
    (args.output_dir / "delisting-import-result.json").write_text(json.dumps(result, indent=2))
    print(json.dumps(result))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
