"""Identify report attachments by content and extract native text/workbooks.

Legacy downloads retain their names for source compatibility. Their extension
is never used as evidence of format. Original attachments are not modified.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import os
from pathlib import Path
import re
import tempfile
import zipfile


PARSER_VERSION = "report-attachments-v1"


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def detect_format(data: bytes) -> tuple[str, str | None]:
    if b"%PDF-" in data[:1024]:
        return "pdf", None
    if data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        return "xls", None
    if data.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "xl/workbook.xml" in archive.namelist():
                return "xlsx", None
        raise ValueError("unsupported ZIP attachment: not an Excel workbook")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        stripped = text.lstrip().lower()
        if stripped.startswith(("<", "{", "[", "access denied", "forbidden")):
            raise ValueError("HTML/JSON/error response is not a report attachment")
        if text.strip() and not any(ord(c) < 32 and c not in "\n\r\t" for c in text):
            return "text", encoding
    raise ValueError("unsupported or corrupt report attachment")


def validate_download(data: bytes) -> str:
    kind, _ = detect_format(data)
    if kind == "pdf":
        import fitz
        with fitz.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass or len(doc) == 0:
                raise ValueError("PDF is encrypted or has no pages")
            for index in range(len(doc)):
                doc.load_page(index)  # Reject a broken page tree before publishing.
    elif kind == "xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            if not book.nsheets:
                raise ValueError("Excel workbook has no worksheets")
        finally:
            book.release_resources()
    elif kind == "xlsx":
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if not book.sheetnames:
                raise ValueError("Excel workbook has no worksheets")
        finally:
            book.close()
    return kind


def safe_stem(stem: str, max_bytes: int = 150) -> str:
    if len(stem.encode("utf-8")) <= max_bytes:
        return stem
    suffix = "_" + hashlib.sha1(stem.encode("utf-8")).hexdigest()[:12]
    prefix = stem.encode("utf-8")[:max_bytes - len(suffix)].decode("utf-8", errors="ignore")
    return prefix.rstrip("._- ") + suffix


def output_paths(path: Path, directory: Path) -> tuple[Path, Path, Path]:
    stem = directory / safe_stem(path.stem)
    return stem.with_suffix(".md"), stem.with_name(stem.name + "_quality.json"), stem.with_name(stem.name + "_quality.md")


def atomic_write(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix="." + path.name + ".", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def _column(index: int) -> str:
    value = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        value = chr(65 + remainder) + value
    return value


def _escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("|", "\\|").replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")


def _number(value: float, number_format: str) -> str:
    # Keep the stored value in the cell manifest as well as its displayed unit.
    if "%" in number_format:
        match = re.search(r"0\.([0#]+)%", number_format)
        digits = len(match.group(1)) if match else 0
        return f"{value * 100:.{digits}f}%"
    return str(int(value)) if value == int(value) else repr(value)


def extract_native(data: bytes) -> tuple[str, dict, list[dict] | None]:
    kind, encoding = detect_format(data)
    if kind == "text":
        text = data.decode(encoding).replace("\r\n", "\n").replace("\r", "\n")
        return text, {"attachment_format": kind, "encoding": encoding, "pages": 1, "extraction_method": "native_text", "warnings": []}, None
    if kind not in {"xls", "xlsx"}:
        raise ValueError("PDF requires the PDF conversion path")
    sheets = []
    if kind == "xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=data, formatting_info=True, on_demand=True)
        try:
            for sheet in book.sheets():
                cells = []
                for row in range(sheet.nrows):
                    for col in range(sheet.ncols):
                        cell = sheet.cell(row, col)
                        if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            continue
                        fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            display = xlrd.xldate_as_datetime(value, book.datemode).isoformat(sep=" ")
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            display = xlrd.error_text_from_code.get(value, f"Excel error {value}")
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            display = "TRUE" if value else "FALSE"
                        elif cell.ctype == xlrd.XL_CELL_NUMBER:
                            display = _number(value, fmt)
                        else:
                            display = str(value)
                        cells.append({"row": row + 1, "column": col + 1, "address": f"{_column(col + 1)}{row + 1}", "type": cell.ctype, "value": value, "display": display, "format": fmt, "is_error": cell.ctype == xlrd.XL_CELL_ERROR})
                sheets.append({"name": sheet.name, "rows": sheet.nrows, "columns": sheet.ncols, "visibility": sheet.visibility, "merged_cells": sheet.merged_cells, "cells": cells})
        finally:
            book.release_resources()
    else:
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                cells = []
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        value = cell.value
                        display = value.isoformat() if isinstance(value, (dt.date, dt.datetime, dt.time)) else _number(value, cell.number_format) if isinstance(value, (int, float)) and not isinstance(value, bool) else str(value)
                        cells.append({"row": cell.row, "column": cell.column, "address": cell.coordinate, "type": cell.data_type, "value": value, "display": display, "format": cell.number_format, "is_error": cell.data_type == "e"})
                sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column, "cells": cells})
        finally:
            book.close()
    if not sheets or not any(s["cells"] for s in sheets):
        raise ValueError("Excel workbook contains no readable cells")
    parts = ["工作簿附件：以下为各工作表保存的单元格值，公式未重新计算。", ""]
    for sheet in sheets:
        parts.extend([f"## 工作表：{sheet['name']}", ""])
        parts.append("| 行 | " + " | ".join(_column(i + 1) for i in range(sheet["columns"])) + " |")
        parts.append("| --- | " + " | ".join("---" for _ in range(sheet["columns"])) + " |")
        rows = {}
        for cell in sheet["cells"]:
            rows.setdefault(cell["row"], {})[cell["column"]] = cell["display"]
        for index, values in rows.items():
            parts.append(f"| {index} | " + " | ".join(_escape(values.get(c + 1, "")) for c in range(sheet["columns"])) + " |")
        parts.append("")
    errors = sum(c["is_error"] for s in sheets for c in s["cells"])
    details = {"attachment_format": kind, "pages": len(sheets), "extraction_method": "native_spreadsheet", "sheet_count": len(sheets), "cell_count": sum(len(s["cells"]) for s in sheets), "source_error_cells": errors, "warnings": ["source_contains_excel_error_values"] if errors else [], "omitted_objects": "Charts/images/macros are retained in the original workbook; this export covers all worksheet cell values, including hidden sheets."}
    return "\n".join(parts), details, sheets


def convert_native(path: Path, directory: Path) -> dict:
    data = path.read_bytes()
    text, details, cells = extract_native(data)
    md, quality, quality_md = output_paths(path, directory)
    summary = {"source_pdf": str(path.resolve()), "parser_version": PARSER_VERSION, "search_full_body": True, "source_sha256": sha256(data), "source_bytes": len(data), "content_sha256": sha256(text.encode()), "ok_pages": details["pages"], "failed_pages": [], "warning_pages": {}, "total_chars": len(text), "page_quality": [], **details}
    if cells is not None:
        manifest = md.with_name(md.stem + "_cells.json")
        atomic_write(manifest, json.dumps(cells, ensure_ascii=False, default=str).encode())
        summary["cells_manifest"] = str(manifest.resolve())
        summary["cells_manifest_sha256"] = sha256(manifest.read_bytes())
    atomic_write(md, text.encode())
    atomic_write(quality_md, (f"# 附件解析结果\n\n格式：{details['attachment_format']}\n\n解析单元：{details['pages']}\n\n字符数：{len(text)}\n\n失败：0\n").encode())
    atomic_write(quality, (json.dumps(summary, ensure_ascii=False, indent=2) + "\n").encode())
    validate_outputs(path, md, quality)
    return summary


def validate_outputs(path: Path, md: Path, quality: Path) -> dict:
    if not md.exists() or not md.stat().st_size or not quality.exists():
        raise ValueError("missing attachment text or quality report")
    summary = json.loads(quality.read_text())
    if summary.get("error") or summary.get("failed_pages"):
        raise ValueError(f"attachment conversion incomplete: {summary.get('error') or summary.get('failed_pages')}")
    if not summary.get("pages") or summary.get("total_chars", 0) <= 0:
        raise ValueError("attachment conversion produced no content")
    if summary.get("source_sha256") and summary["source_sha256"] != sha256(path.read_bytes()):
        raise ValueError("attachment source changed since conversion")
    if summary.get("content_sha256") and summary["content_sha256"] != sha256(md.read_bytes()):
        raise ValueError("attachment output checksum mismatch")
    return summary
